"""Prescription .docx rendering with python-docx.

Layout, top to bottom: doctor header block (details left, photo right), a thin rule, the meta
row, patient block, diagnosis (+ ICD-10 table), medications table, tests, advice, follow-up,
the AI safety flags box (drafts only) and the signature block. Drafts carry a banner at the top
of every page (in the page header) and a draft footer; approved copies carry the approval
footer. A4, 2 cm margins, Calibri.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timedelta, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

logger = logging.getLogger(__name__)

IST = timezone(timedelta(hours=5, minutes=30))
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DASH = "—"
DRAFT_TEXT = "DRAFT: AI-generated, not valid until reviewed and approved by the doctor"

SEVERITY_COLORS = {
    "high": RGBColor(0xB0, 0x00, 0x20),
    "medium": RGBColor(0xA8, 0x6A, 0x00),
    "low": RGBColor(0x55, 0x55, 0x55),
}


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def _rule(doc) -> None:
    """Thin horizontal rule as a paragraph bottom border."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(4)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (
        ("w:val", "single"),
        ("w:sz", "6"),
        ("w:space", "1"),
        ("w:color", "888888"),
    ):
        bottom.set(qn(attr), value)
    border.append(bottom)
    para._p.get_or_add_pPr().append(border)


def _add_field(run, instruction: str) -> None:
    """Insert a Word field (e.g. PAGE) into a run."""
    for kind, text in (
        ("begin", None),
        (None, instruction),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = instruction
        else:
            element = OxmlElement("w:t")
            element.text = text
        run._r.append(element)


def _text(
    paragraph,
    text: str,
    size: float = 10,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _set_widths(table, widths_cm: list[float]) -> None:
    """Fixed column widths (python-docx needs them on every cell for Word and LibreOffice)."""
    table.autofit = False
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)  # writes <w:gridCol>, which LibreOffice honours
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def _cell_text(cell, text: str, size: float = 9, bold: bool = False) -> None:
    cell.text = ""
    _text(cell.paragraphs[0], text, size=size, bold=bold)


def _format_date(value, fmt: str = "%d %b %Y") -> str:
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return value
    return value.astimezone(IST).strftime(fmt) if value.tzinfo else value.strftime(fmt)


def _doctor_header(doc, doctor: dict, photo_bytes: bytes | None) -> None:
    """First block of the document: doctor details on the left, optional photo on the right."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    left, right = table.rows[0].cells
    left.width, right.width = Cm(14.0), Cm(3.0)

    first = left.paragraphs[0]
    _text(
        first,
        doctor.get("full_name") or "Doctor",
        size=17,
        bold=True,
        color=RGBColor(0x14, 0x4E, 0x6E),
    )
    if doctor.get("specialization"):
        _text(
            left.add_paragraph(),
            doctor["specialization"],
            size=11,
            color=RGBColor(0x44, 0x44, 0x44),
        )
    reg = f"Reg. No. {doctor.get('reg_number') or DASH}, {doctor.get('council') or DASH}"
    _text(left.add_paragraph(), reg, size=10, bold=True)
    if doctor.get("verified_at"):
        _text(
            left.add_paragraph(),
            f"Verified on ClinicalScribe ({_format_date(doctor['verified_at'])})",
            size=8.5,
            color=RGBColor(0x1E, 0x6B, 0x35),
        )
    clinic = [doctor.get("clinic_name"), doctor.get("clinic_address")]
    if doctor.get("clinic_phone"):
        clinic.append(f"Phone: {doctor['clinic_phone']}")
    clinic_line = " | ".join(part for part in clinic if part)
    if clinic_line:
        _text(left.add_paragraph(), clinic_line, size=9, color=RGBColor(0x55, 0x55, 0x55))

    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if photo_bytes:
        try:
            right.paragraphs[0].add_run().add_picture(io.BytesIO(photo_bytes), width=Cm(2.5))
        except Exception:
            logger.warning("Could not add the doctor photo to the document; continuing without it")


def build_prescription_docx(
    content: dict,
    flags: list[dict],
    doctor: dict,
    patient: dict,
    is_draft: bool = True,
    approval_code: str | None = None,
    approved_at: datetime | None = None,
    consult_id: str | None = None,
    prescription_id: str | None = None,
    doctor_photo_bytes: bytes | None = None,
) -> bytes:
    """Build a prescription .docx and return its bytes."""
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(3)

    # Draft banner lives in the page header, so it tops every page while the doctor header
    # remains the first block of the document body.
    if is_draft:
        header_para = section.header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _text(header_para, f"⚠ {DRAFT_TEXT}", size=10, bold=True, color=RGBColor(0xB4, 0x5A, 0x00))

    _doctor_header(doc, doctor, doctor_photo_bytes)
    _rule(doc)

    # Meta row
    stamp_source = approved_at or datetime.now(timezone.utc)
    meta_parts = []
    if approval_code:
        meta_parts.append(f"Prescription: {approval_code}")
    elif prescription_id:
        meta_parts.append(f"Prescription ID: {prescription_id}")
    else:
        meta_parts.append("Prescription: DRAFT")
    meta_parts.append(f"Date: {stamp_source.astimezone(IST).strftime('%d %b %Y, %I:%M %p')} IST")
    if consult_id:
        meta_parts.append(f"Consult ID: {consult_id}")
    _text(doc.add_paragraph(), "  |  ".join(meta_parts), size=8, color=RGBColor(0x66, 0x66, 0x66))

    # Patient block
    doc.add_heading("Patient", level=2)
    info = doc.add_table(rows=1, cols=3)
    age, gender = patient.get("age"), patient.get("gender")
    age_gender = (
        " / ".join(str(v) for v in (f"{age} yrs" if age is not None else None, gender) if v) or DASH
    )
    _cell_text(info.rows[0].cells[0], f"Name: {patient.get('full_name') or DASH}")
    _cell_text(info.rows[0].cells[1], f"Age / Gender: {age_gender}")
    _cell_text(info.rows[0].cells[2], f"Patient ID: {patient.get('id') or DASH}", size=8)

    # Diagnosis
    diagnoses = content.get("diagnosis") or []
    if diagnoses:
        doc.add_heading("Diagnosis", level=2)
        for item in diagnoses:
            doc.add_paragraph(item, style="List Bullet")
    icd10 = content.get("icd10") or []
    if icd10:
        _text(doc.add_paragraph(), "ICD-10 suggestions", size=9, italic=True)
        icd = doc.add_table(rows=1, cols=2)
        icd.style = "Light Grid Accent 1"
        _cell_text(icd.rows[0].cells[0], "Code", bold=True)
        _cell_text(icd.rows[0].cells[1], "Description", bold=True)
        for entry in icd10:
            row = icd.add_row().cells
            _cell_text(row[0], str(entry.get("code") or DASH))
            _cell_text(row[1], str(entry.get("description") or DASH))
        _set_widths(icd, [3.0, 14.0])

    # Medications
    medications = content.get("medications") or []
    if medications:
        doc.add_heading("Medications (Rx)", level=2)
        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"
        for i, title in enumerate(
            ("#", "Drug (strength)", "Dose", "Route", "Frequency", "Duration", "Instructions")
        ):
            _cell_text(table.rows[0].cells[i], title, bold=True)
        for number, med in enumerate(medications, 1):
            row = table.add_row().cells
            drug = med.get("drug_name") or DASH
            if med.get("strength"):
                drug = f"{drug} ({med['strength']})"
            values = [
                str(number),
                drug,
                med.get("dose"),
                med.get("route"),
                med.get("frequency"),
                med.get("duration"),
                med.get("instructions"),
            ]
            for cell, value in zip(row, values):
                _cell_text(cell, value or DASH)
        _set_widths(table, [0.8, 3.6, 2.0, 1.7, 2.9, 2.0, 4.0])

    for title, items in (
        ("Tests advised", content.get("tests_advised")),
        ("Advice", content.get("advice")),
    ):
        if items:
            doc.add_heading(title, level=2)
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
    if content.get("follow_up"):
        doc.add_heading("Follow-up", level=2)
        doc.add_paragraph(content["follow_up"])
    if content.get("notes") and is_draft:
        doc.add_heading("Notes", level=2)
        doc.add_paragraph(content["notes"])

    # AI safety flags (drafts only): a shaded box
    if is_draft and flags:
        doc.add_heading("AI safety flags", level=2)
        box = doc.add_table(rows=1, cols=1)
        box.style = "Table Grid"
        cell = box.rows[0].cells[0]
        _shade(cell, "FFF4E5")
        cell.text = ""
        for index, flag in enumerate(flags):
            para = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            severity = flag.get("severity", "low")
            _text(
                para,
                f"[{severity.upper()}] ",
                size=9,
                bold=True,
                color=SEVERITY_COLORS.get(severity),
            )
            _text(para, flag.get("message", ""), size=9)
            if flag.get("field_ref"):
                _text(para, f"  ({flag['field_ref']})", size=8, color=RGBColor(0x77, 0x77, 0x77))

    # Signature block
    _rule(doc)
    if is_draft:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(18)
        _text(para, "_" * 42, size=10)
        _text(doc.add_paragraph(), "Doctor's signature", size=9, bold=True)
    else:
        _text(doc.add_paragraph(), doctor.get("full_name") or "Doctor", size=11, bold=True)
        _text(
            doc.add_paragraph(),
            f"Digitally approved on {stamp_source.astimezone(IST).strftime('%d %b %Y, %I:%M %p')} IST "
            f"({stamp_source.astimezone(timezone.utc).strftime('%H:%M')} UTC)",
            size=9,
        )
        if approval_code:
            _text(doc.add_paragraph(), f"Approval code: {approval_code}", size=9, bold=True)

    # Footer: draft notice or approval line, plus page numbers on every page
    footer = section.footer
    footer.is_linked_to_previous = False
    note = footer.paragraphs[0]
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_draft:
        _text(note, DRAFT_TEXT, size=8, bold=True, color=RGBColor(0xB4, 0x5A, 0x00))
    else:
        _text(
            note,
            f"Generated with ClinicalScribe. Approved by {doctor.get('full_name') or 'Doctor'}.",
            size=8,
            color=RGBColor(0x66, 0x66, 0x66),
        )
    numbers = footer.add_paragraph()
    numbers.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _text(numbers, "Page ", size=8, color=RGBColor(0x66, 0x66, 0x66))
    page = numbers.add_run()
    page.font.size = Pt(8)
    _add_field(page, "PAGE")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
